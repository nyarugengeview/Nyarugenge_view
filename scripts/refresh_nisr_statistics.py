#!/usr/bin/env python3
"""Download and normalize published NISR Nyarugenge population tables.

The script fails closed if its expected source tables or totals change. It never
guesses a value from a graph or replaces an unavailable document with an estimate.
"""
from __future__ import annotations

import json
import re
import csv
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.request import Request, urlopen

import openpyxl
import pdfplumber

ROOT = Path(__file__).resolve().parents[1]
ACQUIRED = ROOT / "data" / "acquired" / "nisr"
OUT = ROOT / "public" / "data" / "statistics"
REPORTS = ROOT / "data" / "reports"
TODAY = date.today().isoformat()
SOURCES = {
    "2022": "https://statistics.gov.rw/sites/default/files/2025-05/Nyarugenge.pdf",
    "2012": "https://statistics.gov.rw/sites/default/files/2025-06/Nyarugenge_RPHC4_District_Profile.pdf",
    "projections": "https://statistics.gov.rw/sites/default/files/2025-06/1.1%20Nyarugenge_single_age.xlsx",
    "gender": "https://statistics.gov.rw/sites/default/files/2025-08/Gender%20Statistics%20Profile%20Report%20Nyarugenge.pdf",
}
SECTORS_2022 = ["Gitega", "Kanyinya", "Kigali", "Kimisagara", "Mageregere", "Muhima", "Nyakabanda", "Nyamirambo", "Nyarugenge", "Rwezamenyo"]
SECTORS_2012 = ["Gitega", "Kanyinya", "Kigali", "Kimisagara", "Mageragere", "Muhima", "Nyakabanda", "Nyamirambo", "Nyarugenge", "Rwezamenyo"]


def fetch(url: str, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    request = Request(url, headers={"User-Agent": "NyarugengeViewDataPipeline/1.0"})
    with urlopen(request, timeout=90) as response:
        if response.status != 200:
            raise RuntimeError(f"Download failed ({response.status}): {url}")
        payload = response.read()
        if len(payload) < 5000:
            raise RuntimeError(f"Downloaded file is unexpectedly small ({len(payload)} bytes): {url}")
        path.write_bytes(payload)


def pdf_text(path: Path, page_numbers: list[int]) -> str:
    with pdfplumber.open(path) as pdf:
        if max(page_numbers) > len(pdf.pages):
            raise RuntimeError(f"Expected pages absent in {path.name} (has {len(pdf.pages)} pages).")
        return "\n".join(pdf.pages[p - 1].extract_text() or "" for p in page_numbers)


def parse_rows(text: str, names: list[str], pattern: str, labels: list[str]) -> list[dict]:
    regex = re.compile(pattern)
    rows = []
    for line in text.splitlines():
        match = regex.match(line.strip())
        if not match:
            continue
        name = match.group(1)
        if name not in names:
            continue
        values = match.groups()[1:]
        rows.append({"sector": name, **{key: value for key, value in zip(labels, values)}})
    seen = {row["sector"] for row in rows}
    if seen != set(names) or len(rows) != len(names):
        raise RuntimeError(f"Could not parse every expected sector. Found: {sorted(seen)}; expected: {names}")
    return rows


def num(value: str | None) -> int | float | None:
    if value is None or value in {"", "-", "—"}:
        return None
    return float(value.replace(",", "")) if "." in value else int(value.replace(",", ""))


def parse_2022(pdf_path: Path) -> dict:
    sex_text = pdf_text(pdf_path, [17])
    residence_text = pdf_text(pdf_path, [18])
    sex = parse_rows(
        sex_text, SECTORS_2022,
        r"^(.+?)\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)\s+100\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)$",
        ["population", "male", "female", "male_share_pct", "female_share_pct", "district_population_share_pct"],
    )
    residence = parse_rows(
        residence_text, SECTORS_2022,
        r"^(.+?)\s+([\d,]+)\s+([\d,]+|-)\s+([\d,]+|-)\s+100\s+([\d.]+|-)\s+([\d.]+|-)$",
        ["population", "urban", "rural", "urban_share_pct", "rural_share_pct"],
    )
    by_residence = {row["sector"]: row for row in residence}
    result = []
    for sex_row in sex:
        r = by_residence[sex_row["sector"]]
        result.append({
            "sector": sex_row["sector"],
            "population": int(num(sex_row["population"])),
            "male": int(num(sex_row["male"])),
            "female": int(num(sex_row["female"])),
            "maleSharePct": float(sex_row["male_share_pct"]),
            "femaleSharePct": float(sex_row["female_share_pct"]),
            "districtPopulationSharePct": float(sex_row["district_population_share_pct"]),
            "urban": num(r["urban"]),
            "rural": num(r["rural"]),
            "urbanSharePct": num(r["urban_share_pct"]),
            "ruralSharePct": num(r["rural_share_pct"]),
        })
    totals = {key: sum(row[key] or 0 for row in result) for key in ["population", "male", "female", "urban", "rural"]}
    expected = {"population": 374319, "male": 195780, "female": 178539, "urban": 315899, "rural": 58420}
    if totals != expected:
        raise RuntimeError(f"2022 sector totals do not reconcile to NISR district values: {totals}")
    return {"rows": result, "totals": totals}


def parse_2012(pdf_path: Path) -> dict:
    text = pdf_text(pdf_path, [18, 20])
    sector_text = text.split("Table 2:", 1)[1].split("Source: Rwanda 4th", 1)[0]
    residence_text = text.split("Table 3:", 1)[1].split("Source: Rwanda 4th", 1)[0]
    sex = parse_rows(
        sector_text, SECTORS_2012,
        r"^(.+?)\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)\s+([\d.]+)\s+([\d.]+)\s+([\d,]+)$",
        ["population", "male", "female", "female_share_pct", "district_population_share_pct", "density_per_km2"],
    )
    residence = parse_rows(
        residence_text, SECTORS_2012,
        r"^(.+?)\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)\s+100\s+([\d.]+)\s+([\d.]+)$",
        ["population", "urban", "rural", "urban_share_pct", "rural_share_pct"],
    )
    by_residence = {row["sector"]: row for row in residence}
    result = []
    for row in sex:
        r = by_residence[row["sector"]]
        result.append({
            "sector": row["sector"],
            "population": int(num(row["population"])),
            "male": int(num(row["male"])),
            "female": int(num(row["female"])),
            "femaleSharePct": float(row["female_share_pct"]),
            "districtPopulationSharePct": float(row["district_population_share_pct"]),
            "densityPerKm2": int(num(row["density_per_km2"])),
            "urban": num(r["urban"]),
            "rural": num(r["rural"]),
            "urbanSharePct": float(r["urban_share_pct"]),
            "ruralSharePct": float(r["rural_share_pct"]),
        })
    totals = {key: sum(row[key] or 0 for row in result) for key in ["population", "male", "female", "urban", "rural"]}
    expected = {"population": 284561, "male": 148132, "female": 136429, "urban": 214020, "rural": 70541}
    if totals != expected:
        raise RuntimeError(f"2012 sector totals do not reconcile to NISR district values: {totals}")
    return {"rows": result, "totals": totals}


def parse_projections(path: Path) -> list[dict]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    header = next(sheet.iter_rows(min_row=4, max_row=5, values_only=True))[0:0]  # Keep the source layout below explicit.
    _ = header
    year_row = next(sheet.iter_rows(min_row=4, max_row=4, values_only=True))
    data_rows = list(sheet.iter_rows(min_row=5, max_row=6, values_only=True))
    headers = next(sheet.iter_rows(min_row=5, max_row=5, values_only=True))
    total_row = next(sheet.iter_rows(min_row=6, max_row=6, values_only=True))
    projections = []
    for col in range(1, 31, 3):
        year = year_row[col]
        total = total_row[col]
        male = total_row[col + 1]
        female = total_row[col + 2]
        if not isinstance(year, int) or year < 2023 or not all(isinstance(v, (int, float)) for v in [total, male, female]) or total != male + female:
            raise RuntimeError(f"Projection workbook structure or totals changed around column {col}.")
        projections.append({"year": year, "total": int(total), "male": int(male), "female": int(female)})
    if len(projections) != 10 or projections[0]["total"] != 383348 or projections[-1]["total"] != 477324:
        raise RuntimeError("Projection series did not match expected years 2023–2032.")
    return projections


def parse_gender_profile(path: Path) -> list[dict]:
    """Extract Nyarugenge-specific gender tables with the original universes and units."""
    with pdfplumber.open(path) as pdf:
        age_tables = pdf.pages[17].extract_tables()
        marital_tables = pdf.pages[19].extract_tables()
    age_table = next((t for t in age_tables if t and any((r[0] or '').strip() == '<1' for r in t if r)), None)
    marital_table = next((t for t in marital_tables if t and any((r[0] or '').strip() == 'Never married' for r in t if r)), None)
    marriage_table = next((t for t in marital_tables if t and any((r[0] or '').strip() == 'Nyarugenge District' for r in t if r)), None)
    if not age_table or not marital_table or not marriage_table:
        raise RuntimeError('Could not locate expected machine-readable tables in the Nyarugenge gender profile.')
    rows = []
    for row in age_table:
        if len(row) < 7 or not row[0]:
            continue
        label = row[0].replace('\n', ' ').strip()
        if label not in {'Total Resident Population','<1','<5','1-4','3-5','3-6','0-14','0-15','0-17','6-11','7-12','12-17','13-18','14-35','16-30','15-24','15-49','15-64','16-64','16-59','14+','16+','18+','21+','60+','65+'}:
            continue
        vals = [re.sub(r'[^\d.]', '', str(v or '').replace(',', '')) for v in row[1:7]]
        if len(vals) != 6 or any(not v for v in vals):
            raise RuntimeError(f'Unexpected age/sex table row: {row}')
        rows.append({'table':'Target age groups','indicator':label,'male':int(float(vals[1])),'female':int(float(vals[2])),'total':int(float(vals[0])),'male_share_pct':float(vals[4]),'female_share_pct':float(vals[5]),'total_share_pct':float(vals[3]),'unit':'residents and percent of total district population','universe':'Nyarugenge resident population, 2022 census'})
    for row in marital_table:
        if len(row) < 4 or not row[0] or row[0].strip() == 'Marital status':
            continue
        label = row[0].replace('\n',' ').strip()
        values=[str(v or '').strip() for v in row[1:4]]
        if label not in {'Never married','Married to one wife/husband officially','Married to one wife/husband not officially','Living in polygamous union','Divorced','Separated','Widowed'}:
            continue
        if len(values)!=3 or any(not re.fullmatch(r'\d+(?:\.\d+)?',v) for v in values):
            raise RuntimeError(f'Unexpected marital-status row: {row}')
        rows.append({'table':'Marital status','indicator':label,'male':None,'female':None,'total':None,'male_share_pct':float(values[0]),'female_share_pct':float(values[1]),'total_share_pct':float(values[2]),'unit':'percent','universe':'Nyarugenge resident population aged 12+, 2022 census'})
    for row in marriage_table:
        if len(row) < 3 or (row[0] or '').strip() != 'Nyarugenge District':
            continue
        rows.append({'table':'Mean age at first marriage','indicator':'Nyarugenge District','male':None,'female':None,'total':None,'male_share_pct':float(row[1]),'female_share_pct':float(row[2]),'total_share_pct':None,'unit':'years','universe':'Residents who marry before age 50; Nyarugenge district, 2022 census'})
    if len(rows) != 34 or not any(r['table']=='Mean age at first marriage' and r['male_share_pct']==31.4 and r['female_share_pct']==26.8 for r in rows):
        raise RuntimeError(f'Gender profile extraction did not pass expected row and value checks ({len(rows)} rows).')
    return rows


def write_json(name: str, payload: dict) -> None:
    path = OUT / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> None:
    files = {
        "2022": ACQUIRED / "nyarugenge-2022-rphc.pdf",
        "2012": ACQUIRED / "nyarugenge-2012-rphc4.pdf",
        "projections": ACQUIRED / "nyarugenge-population-projections-2023-2032.xlsx",
        "gender": ACQUIRED / "nyarugenge-gender-profile-2023.pdf",
    }
    for key, path in files.items():
        fetch(SOURCES[key], path)
        print(f"Retrieved {key}: {path.stat().st_size:,} bytes")
    population_2022 = parse_2022(files["2022"])
    population_2012 = parse_2012(files["2012"])
    projections = parse_projections(files["projections"])
    gender_rows = parse_gender_profile(files["gender"])
    retrieved = {
        "retrievedAt": TODAY,
        "sourceOrganization": "National Institute of Statistics of Rwanda (NISR)",
        "sourceUrl": "https://statistics.gov.rw/district-statistics/kigali-city",
    }
    write_json("population-sectors-2022.json", {
        "datasetId": "nisr-phc5-nyarugenge-sectors-2022",
        "title": "Resident population by sector and area of residence",
        "referencePeriod": "2022 census",
        "geography": {"name": "Nyarugenge", "level": "Sector"},
        "sourceDocument": {"title": "The Fifth Rwanda Population and Housing Census, District Profile: Nyarugenge", "publicationDate": "September 2023", "sourceUrl": SOURCES["2022"], "tables": ["Table 1.2, printed p. 5 / PDF p. 17", "Table 1.3, printed p. 6 / PDF p. 18"]},
        "units": {"population": "residents", "shares": "percent", "area": "inhabitants per square kilometre"},
        "rows": population_2022["rows"], "districtTotals": population_2022["totals"],
        "notes": ["Census observations, not current estimates.", "Sector names and district geography follow the 2022 NISR census profile.", "A dash in the published residence table is preserved as null; it is not silently rewritten as a published numeric zero."],
        **retrieved,
    })
    write_json("population-sectors-2012.json", {
        "datasetId": "nisr-phc4-nyarugenge-sectors-2012",
        "title": "Resident population by sector and area of residence",
        "referencePeriod": "2012 census",
        "geography": {"name": "Nyarugenge", "level": "Sector"},
        "sourceDocument": {"title": "Fourth Population and Housing Census, Rwanda, 2012 District Profile: Nyarugenge", "sourceUrl": "https://statistics.gov.rw/sites/default/files/2025-06/Nyarugenge_RPHC4_District_Profile.pdf", "tables": ["Table 2, printed p. 3 / PDF p. 18", "Table 3, printed p. 5 / PDF p. 20"]},
        "units": {"population": "residents", "shares": "percent", "densityPerKm2": "inhabitants per square kilometre"},
        "rows": population_2012["rows"], "districtTotals": population_2012["totals"],
        "notes": ["Census observations, not current estimates.", "2012 and 2022 sector labels/geographies are not assumed to be fully comparable; displayed as separate census snapshots, with no growth rate calculated."],
        **retrieved,
    })
    write_json("population-projections-2023-2032.json", {
        "datasetId": "nisr-nyarugenge-population-projections-2023-2032",
        "title": "Projected resident population by sex",
        "referencePeriod": "2023–2032 projections",
        "geography": {"name": "Nyarugenge", "level": "District"},
        "sourceDocument": {"title": "Subnational Population Projections (2023–2032): Nyarugenge", "sourceUrl": SOURCES["projections"], "sheet": "Nyarugenge", "table": "Distribution of the Resident Population by single age and by sex"},
        "units": "residents",
        "series": projections,
        "notes": ["Projected values, not observed census counts.", "The source workbook cites the Rwanda Population and Housing Census (NISR 2022) as its source."],
        **retrieved,
    })
    gender_path = OUT / "population-gender-profile-2022.csv"
    gender_path.parent.mkdir(parents=True, exist_ok=True)
    with gender_path.open('w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=['table','indicator','male','female','total','male_share_pct','female_share_pct','total_share_pct','unit','universe'])
        writer.writeheader(); writer.writerows(gender_rows)
    REPORTS.mkdir(parents=True, exist_ok=True)
    inventory = {
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "sources": [
            {"title": "Nyarugenge Population and Housing Census 2022 District Profile", "url": SOURCES["2022"], "access": "downloaded; PDF text tables machine-extracted", "referencePeriod": "2022 census", "outputs": ["population-sectors-2022.json"], "license": "No explicit reuse license found on the attachment page; linked to NISR source and attributed."},
            {"title": "Nyarugenge Population and Housing Census 2012 District Profile", "url": SOURCES["2012"], "access": "downloaded; PDF text tables machine-extracted", "referencePeriod": "2012 census", "outputs": ["population-sectors-2012.json"], "license": "No explicit reuse license found on the attachment page; linked to NISR source and attributed."},
            {"title": "Nyarugenge Subnational Population Projections 2023–2032", "url": SOURCES["projections"], "access": "downloaded; XLSX values extracted", "referencePeriod": "2023–2032", "outputs": ["population-projections-2023-2032.json"], "license": "No explicit reuse license found on the attachment page; linked to NISR source and attributed."},
            {"title": "Gender Statistics Profile: Nyarugenge", "url": SOURCES["gender"], "access": "downloaded; machine-readable tables extracted and checked", "referencePeriod": "2022 census tables", "outputs": ["population-gender-profile-2022.csv"], "tables": ["Table 1, printed p. 17 / PDF p. 18", "Table 2, printed p. 19 / PDF p. 20", "Table 3, printed p. 19 / PDF p. 20"], "license": "No explicit reuse license found on the attachment page; linked to NISR source and attributed."},
        ],
        "validation": {"2022SectorTotals": population_2022["totals"], "2012SectorTotals": population_2012["totals"], "projectionYears": [projections[0]["year"], projections[-1]["year"]], "projectionTotals": [projections[0]["total"], projections[-1]["total"]], "genderProfileRows": len(gender_rows)},
    }
    (REPORTS / "nisr-statistics-inventory.json").write_text(json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf-8")
    print("Validated census table sums, projection series, and Nyarugenge gender tables; wrote normalized datasets.")


if __name__ == "__main__":
    main()

